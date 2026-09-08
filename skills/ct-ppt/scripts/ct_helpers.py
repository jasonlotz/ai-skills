"""
Cleartelligence PowerPoint helpers.
Import these in any build script:
    from scripts.ct_helpers import add_slide, set_ph, fill_body, inspect_slide
"""

from pptx.dml.color import RGBColor

# ── Brand colors ───────────────────────────────────────────────────
NAVY    = RGBColor(0x18, 0x2D, 0x4F)   # Primary
BLUE    = RGBColor(0x33, 0x64, 0xAC)   # Secondary
LT_BLUE = RGBColor(0x3D, 0x8E, 0xC8)   # Accent
ORANGE  = RGBColor(0xEC, 0x8E, 0x2E)   # Highlight
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x11, 0x0B, 0x18)   # Body text


def inspect_slide(slide):
    """Print all placeholder indices and types on a slide. Useful for debugging layouts."""
    for ph in slide.placeholders:
        fmt = ph.placeholder_format
        print(f"  idx={fmt.idx}  type={fmt.type}  name='{ph.name}'")


def list_layouts(prs):
    """Print every layout across ALL masters with its placeholder indices.

    The template has three masters. prs.slide_layouts only exposes the first,
    which is why layouts like 'Divider Solid' and 'End Slide' appear missing.
    Always locate layouts by name with add_slide(), not by index.
    """
    for mi, m in enumerate(prs.slide_masters):
        print(f"=== master {mi} ===")
        for l in m.slide_layouts:
            idxs = ",".join(str(ph.placeholder_format.idx) for ph in l.placeholders)
            print(f"  {l.name:<28} ph idx: {idxs or '-'}")


def find_layout(prs, name, master=None):
    """Find a layout by exact name across all masters. master pins the search."""
    masters = ([prs.slide_masters[master]] if master is not None
               else list(prs.slide_masters))
    for m in masters:
        for l in m.slide_layouts:
            if l.name == name:
                return l
    raise KeyError(
        f"layout {name!r} not found. Available: "
        + ", ".join(sorted({l.name for m in prs.slide_masters
                            for l in m.slide_layouts})))


def add_slide(prs, layout, master=None):
    """Add a slide. `layout` is a layout NAME (preferred) or a 1-based index
    into master 0 for backwards compatibility."""
    if isinstance(layout, int):
        return prs.slides.add_slide(prs.slide_layouts[layout - 1])
    return prs.slides.add_slide(find_layout(prs, layout, master))


def add_textbox(slide, left_in, top_in, width_in, height_in, text,
                size=18, bold=False, color=None, align=None):
    """Add a free text box. This template carries most content in text boxes
    rather than placeholders, so this is the main content tool."""
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in),
                                   Inches(width_in), Inches(height_in))
    tf = box.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, (list, tuple)) else [text]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = str(line)
        para.font.size = Pt(size)
        para.font.bold = bold
        para.font.color.rgb = color if color is not None else DARK
        if align == "center":
            para.alignment = PP_ALIGN.CENTER
        elif align == "right":
            para.alignment = PP_ALIGN.RIGHT
    return box


def add_table(slide, left_in, top_in, width_in, height_in, headers, rows,
              font_size=12):
    """Add a branded table: navy header row, alternating light rows."""
    from pptx.util import Inches, Pt
    shape = slide.shapes.add_table(len(rows) + 1, len(headers),
                                   Inches(left_in), Inches(top_in),
                                   Inches(width_in), Inches(height_in))
    tbl = shape.table
    for c, h in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = str(h)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        para = cell.text_frame.paragraphs[0]
        para.font.bold = True; para.font.size = Pt(font_size)
        para.font.color.rgb = WHITE
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (RGBColor(0xF4, 0xF5, 0xF7) if r % 2
                                        else WHITE)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(font_size); para.font.color.rgb = DARK
    return tbl


def set_ph(slide, idx, text):
    """Set a placeholder by its idx. Does nothing if idx not present."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            return ph
    return None


def fill_body(slide, idx, bullets, bold_first=False):
    """
    Fill a body placeholder with bullet points.
    bullets: list of strings
    bold_first: if True, the first bullet is rendered bold at level 0
    """
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            tf = ph.text_frame
            tf.clear()
            tf.word_wrap = True
            for i, line in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.level = 0
                if bold_first and i == 0:
                    p.font.bold = True
            return ph
    return None


_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def clear_sample_slides(prs):
    """Remove all existing slides from the template before adding new content.

    Drops each slide's RELATIONSHIP as well as its entry in the slide-id list.
    Removing only the id-list entry leaves the slide parts orphaned inside the
    package; new slides then reuse those part names and the saved .pptx ends up
    with duplicate zip members, which PowerPoint reports as a file needing
    repair. Dropping the rel lets the part be garbage-collected on save.
    """
    id_list = prs.slides._sldIdLst
    for sld in list(id_list):
        rId = sld.get(_R_ID)
        if rId:
            try:
                prs.part.drop_rel(rId)
            except KeyError:
                pass
        id_list.remove(sld)
